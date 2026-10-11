"""Excel extraction keeps the sheet name and the header of every cell."""
import io

from openpyxl import Workbook

import ingest


def write(tmp_path, build):
    wb = Workbook()
    build(wb)
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return str(path)


def test_cells_are_header_value_pairs_with_the_sheet_name(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Payroll"
        ws.append(["Name", "SSN", "Routing", "Account"])
        ws.append(["Jane Doe", "123-45-6789", "021000021", "4471029381"])
        ws.append(["Sam Roe", "", "", "88120034"])  # empty cells are skipped

    text = ingest.extract_text_from_excel(write(tmp_path, build))

    assert text.splitlines() == [
        "Payroll row 2: Name: Jane Doe | SSN: 123-45-6789 | Routing: 021000021 | Account: 4471029381",
        "Payroll row 3: Name: Sam Roe | Account: 88120034",
    ]


def test_every_sheet_is_included_and_header_only_sheets_are_text(tmp_path):
    def build(wb):
        wb.active.title = "Staff"
        wb.active.append(["Name", "Salary"])
        wb.active.append(["Ana", "98500"])
        notes = wb.create_sheet("Banking")
        notes.append(["Routing number 021000021"])

    lines = ingest.extract_text_from_excel(write(tmp_path, build)).splitlines()

    assert lines == ["Staff row 2: Name: Ana | Salary: 98500", "Banking: Routing number 021000021"]


def test_sheet_without_a_header_row_uses_column_letters(tmp_path):
    def build(wb):
        wb.active.title = "Raw"
        wb.active.append(["123-45-6789", "98500"])
        wb.active.append(["234-56-7890", "77000"])

    lines = ingest.extract_text_from_excel(write(tmp_path, build)).splitlines()

    assert lines[0] == "Raw row 1: Column A: 123-45-6789 | Column B: 98500"


def test_missing_header_cells_and_extra_columns_get_column_letters(tmp_path):
    def build(wb):
        wb.active.title = "S"
        wb.active.append(["Name", None, "Notes"])
        wb.active.append(["Ana", "x", "y", "z"])

    assert ingest.extract_text_from_excel(write(tmp_path, build)) == "S row 2: Name: Ana | Column B: x | Notes: y | Column D: z"


def test_header_context_lets_the_rules_find_bare_numbers(client, fake_llm):
    """Without "Account:" in front, a bare 10-digit cell is not recognizably an account number."""
    def build(wb):
        wb.active.append(["Employee", "Account"])
        wb.active.append(["Ana", "4471029381"])

    buf = io.BytesIO()
    wb = Workbook()
    build(wb)
    wb.save(buf)
    res = client.post("/upload", files={"file": ("b.xlsx", buf.getvalue(), "application/octet-stream")})

    assert [(f["category"], f["text_snippet"]) for f in res.json()["findings"]] == [("bank_account", "******9381")]
